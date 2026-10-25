"""Extract plain text from uploaded resume files."""
import io

import docx
import pdfplumber


class UnsupportedFileError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return _from_pdf(data)
    if name.endswith(".docx"):
        return _from_docx(data)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise UnsupportedFileError("Only .pdf, .docx and .txt files are supported.")


def _from_pdf(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages)


def _from_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)
