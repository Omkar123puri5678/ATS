"""Generate the sample resumes in samples/resumes/ (PDF, DOCX and TXT).

All people, companies and contact details are fictional. Each resume is
written to land in a different score range so you can see how the analyzer
reacts to strong, average and weak resumes.

    pip install fpdf2 python-docx
    python samples/generate.py
"""
from pathlib import Path

import docx
from fpdf import FPDF

OUT = Path(__file__).parent / "resumes"

# (file name, formats, header lines, [(section, lines)])
RESUMES = [
    (
        "01_strong_software_engineer",
        ["pdf", "docx", "txt"],
        ["Aarav Mehta", "aarav.mehta@example.com | +91 90000 11111 | linkedin.com/in/aarav-mehta-demo | github.com/aarav-demo"],
        [
            ("Summary", [
                "Backend-focused software engineer with 3 years of experience building scalable REST APIs "
                "and cloud services in Python and Node.js.",
            ]),
            ("Experience", [
                "Software Engineer, Brightline Tech (Jul 2023 - Present)",
                "- Designed and built 12 REST APIs with FastAPI and PostgreSQL serving 40,000 daily users",
                "- Reduced p95 API latency by 45% by adding Redis caching and query indexing",
                "- Led migration of 8 services to Docker and Kubernetes on AWS, cutting hosting cost by 30%",
                "- Automated CI/CD with GitHub Actions, reducing release time from 2 days to 3 hours",
                "Software Engineering Intern, Nimbus Labs (Jan 2023 - Jun 2023)",
                "- Developed a React and TypeScript admin dashboard used by 150 support staff",
                "- Wrote 200+ unit tests with Pytest and Jest, raising coverage from 52% to 88%",
            ]),
            ("Projects", [
                "- Built an open-source job tracker in React and Node.js with 1,200 GitHub stars",
                "- Implemented a real-time chat service with WebSockets handling 5,000 concurrent users",
            ]),
            ("Education", ["B.Tech in Computer Science, Demo Institute of Technology, 2019 - 2023, CGPA 8.9"]),
            ("Skills", [
                "Python, JavaScript, TypeScript, SQL, FastAPI, Django, Node.js, React, PostgreSQL, MongoDB, "
                "Redis, Docker, Kubernetes, AWS, Git, CI/CD, Microservices, REST APIs, Linux",
            ]),
            ("Certifications", ["AWS Certified Developer - Associate (2024)"]),
        ],
    ),
    (
        "02_average_data_analyst",
        ["pdf", "docx", "txt"],
        ["Sneha Kulkarni", "sneha.k@example.com | +91 90000 22222"],
        [
            ("Experience", [
                "Data Analyst, Riverstone Retail (2022 - Present)",
                "- Responsible for creating weekly sales reports in Excel",
                "- Worked on dashboards in Power BI for the sales team",
                "- Helped clean customer data using SQL",
                "- Created a Tableau dashboard that cut reporting time by 20%",
                "Junior Analyst, Bluepeak Services (2021 - 2022)",
                "- Was part of a team that analyzed survey data",
                "- Prepared presentations for management",
            ]),
            ("Education", ["B.Sc in Statistics, Demo University, 2018 - 2021"]),
            ("Skills", ["Excel, SQL, Power BI, Tableau, Python (basic), Communication"]),
        ],
    ),
    (
        "03_weak_no_structure",
        ["pdf", "txt"],
        ["Rahul"],
        [
            ("", [
                "I am a hardworking and passionate person looking for a good job in a reputed company where I "
                "can use my skills and grow.",
                "I have worked in marketing for some time. I was responsible for social media and handling "
                "some campaigns. I also did other tasks as given by my manager.",
                "I know MS Office and I am good at talking to people. I am a quick learner and team player.",
                "Hobbies: cricket, music, travelling.",
            ]),
        ],
    ),
    (
        "04_fresher_student",
        ["pdf", "docx", "txt"],
        ["Priya Nair", "priya.nair@example.com | +91 90000 44444 | github.com/priya-demo"],
        [
            ("Summary", ["Final-year computer science student interested in web development and machine learning."]),
            ("Education", ["B.E. in Computer Engineering, Demo College of Engineering, 2022 - 2026, CGPA 8.4"]),
            ("Projects", [
                "- Built a movie recommendation system in Python with scikit-learn and Pandas (85% precision on test set)",
                "- Developed a college event portal in React and Firebase used by 600 students",
                "- Created a REST API in Flask for a library management system",
            ]),
            ("Experience", [
                "Web Development Intern, Pixel Forge (May 2025 - Jul 2025)",
                "- Implemented 6 responsive pages in React and Tailwind CSS",
                "- Fixed 25 UI bugs reported by QA",
            ]),
            ("Skills", ["Python, Java, JavaScript, React, Flask, HTML, CSS, SQL, MySQL, Git, Pandas, scikit-learn"]),
            ("Certifications", ["Machine Learning Specialization (online course, 2025)"]),
        ],
    ),
    (
        "05_digital_marketing",
        ["pdf", "docx", "txt"],
        ["Kabir Shah", "kabir.shah@example.com | +91 90000 55555 | linkedin.com/in/kabir-shah-demo"],
        [
            ("Summary", ["Digital marketer with 4 years of experience in SEO, paid ads and content strategy for D2C brands."]),
            ("Experience", [
                "Digital Marketing Manager, Urban Leaf (2023 - Present)",
                "- Grew organic traffic by 120% in 12 months through SEO and content marketing",
                "- Managed a monthly Google Ads and Meta Ads budget of 15 lakh with 4.2x ROAS",
                "- Launched email campaigns that increased repeat purchases by 18%",
                "Marketing Executive, Coastline Media (2021 - 2023)",
                "- Ran social media for 5 brands, growing combined followers from 20K to 95K",
            ]),
            ("Education", ["MBA in Marketing, Demo School of Business, 2019 - 2021"]),
            ("Skills", ["SEO, SEM, Google Ads, Meta Ads, Google Analytics, Content Marketing, Email Marketing, Canva, Excel"]),
        ],
    ),
]

JOBS = {
    "software_engineer": """Backend Software Engineer
We are hiring a backend engineer with 2+ years of experience.
Requirements: Python, FastAPI or Django, PostgreSQL, Redis, Docker, Kubernetes, AWS, REST APIs,
microservices, CI/CD, Git, unit testing. Experience with React or TypeScript is a plus.""",
    "data_analyst": """Data Analyst
Looking for a data analyst to build dashboards and reports for business teams.
Requirements: SQL, Excel, Power BI or Tableau, Python with Pandas, statistics, data cleaning,
data visualization, stakeholder communication. Experience with A/B testing is a plus.""",
    "digital_marketing": """Digital Marketing Specialist
We need a marketer to grow our online presence.
Requirements: SEO, SEM, Google Ads, Meta Ads, Google Analytics, content marketing, email marketing,
social media marketing, campaign reporting, A/B testing, Excel.""",
    "fresher_web_developer": """Junior Web Developer (Freshers welcome)
Requirements: JavaScript, React, HTML, CSS, Git, REST APIs, basic SQL.
Nice to have: Node.js, Tailwind CSS, Python, any internship or project experience.""",
}


def to_text(header, sections):
    lines = list(header) + [""]
    for title, body in sections:
        if title:
            lines.append(title.upper())
        lines.extend(body)
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def write_pdf(path, header, sections):
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 9, header[0], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9.5)
    for line in header[1:]:
        pdf.multi_cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    for title, body in sections:
        pdf.ln(3)
        if title:
            pdf.set_font("Helvetica", "B", 11.5)
            pdf.cell(0, 7, title.upper(), new_x="LMARGIN", new_y="NEXT")
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(1.5)
        pdf.set_font("Helvetica", size=10)
        for line in body:
            pdf.multi_cell(0, 5.2, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def write_docx(path, header, sections):
    d = docx.Document()
    d.add_heading(header[0], level=0)
    for line in header[1:]:
        d.add_paragraph(line)
    for title, body in sections:
        if title:
            d.add_heading(title, level=1)
        for line in body:
            if line.startswith("- "):
                d.add_paragraph(line[2:], style="List Bullet")
            else:
                d.add_paragraph(line)
    d.save(str(path))


def main():
    OUT.mkdir(exist_ok=True)
    jobs_dir = Path(__file__).parent / "job_descriptions"
    jobs_dir.mkdir(exist_ok=True)
    for name, formats, header, sections in RESUMES:
        if "txt" in formats:
            (OUT / f"{name}.txt").write_text(to_text(header, sections), encoding="utf-8")
        if "pdf" in formats:
            write_pdf(OUT / f"{name}.pdf", header, sections)
        if "docx" in formats:
            write_docx(OUT / f"{name}.docx", header, sections)
    for name, text in JOBS.items():
        (jobs_dir / f"{name}.txt").write_text(text + "\n", encoding="utf-8")
    print(f"Wrote {len(list(OUT.iterdir()))} resumes and {len(JOBS)} job descriptions")


if __name__ == "__main__":
    main()
